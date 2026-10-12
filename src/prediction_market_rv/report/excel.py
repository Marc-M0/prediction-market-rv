"""Excel report: Arbitrage / All matches / Not matched / Unpaired events / Notes sheets."""
from __future__ import annotations

from pathlib import Path

from prediction_market_rv.config import CAPITAL
from prediction_market_rv.core.models import Market


def _fmt_px(bid: float | None, ask: float | None) -> str:
    return "n/a" if bid is None or ask is None else f"{bid:.3f} / {ask:.3f}"


def _yes_no_prices(m: Market) -> tuple[str, str]:
    """('YES bid / ask', 'NO bid / ask'). NO bid = 1 - YES ask, NO ask = 1 - YES bid."""
    yes = _fmt_px(m.yes_bid, m.yes_ask)
    no = "n/a" if m.yes_bid is None or m.yes_ask is None else _fmt_px(1 - m.yes_ask, 1 - m.yes_bid)
    return yes, no


def build_row(arb: dict, p: Market, k: Market, score: float, inverted: bool, flags: list[str]) -> dict:
    """Turn one opportunity into a flat report row, with the position sizing as % of capital."""
    poly_is_first = arb["strategy"].startswith("YES Polymarket")
    if poly_is_first:                       # YES Polymarket + NO Kalshi
        poly_px, kal_px = arb["leg1"], arb["leg2"]
        poly_view_side, kal_side = "YES", "NO"
    else:                                   # YES Kalshi + NO Polymarket
        kal_px, poly_px = arb["leg1"], arb["leg2"]
        poly_view_side, kal_side = "NO", "YES"
    # if the Polymarket question is phrased as a negation, its YES is the other side's NO
    poly_side = poly_view_side if not inverted else ("NO" if poly_view_side == "YES" else "YES")

    poly_cost = poly_px + arb["poly_fee"]            # price + fee, per contract
    kal_cost = kal_px + arb["kalshi_fee"]
    total = poly_cost + kal_cost                      # capital needed per $1 of guaranteed payout
    net = arb["net_edge"]

    p_yes, p_no = _yes_no_prices(p)
    k_yes, k_no = _yes_no_prices(k)
    ok = not any(f == "CHECK-RULES" or f.startswith("WARN") or f.endswith("wide-spread")
                 or f.endswith("no-book") or f == "dust-leg" for f in flags)       # [GUARD] dust-leg added
    return {
        "status": "OK" if ok else "CHECK: " + ", ".join(flags),
        "net_per_unit": net,                          # $ profit per $1 payout (after fees)
        "roi": net / total,                           # return on capital deployed
        "profit_on_capital": CAPITAL * net / total,
        "poly_pct": poly_cost / total,
        "kal_pct": kal_cost / total,
        "action": f"Buy {poly_side} on Polymarket @ {poly_px:.3f}  +  Buy {kal_side} on Kalshi @ {kal_px:.3f}",
        "k_q": k.title, "k_url": k.url, "k_ticker": k.market_id,
        "k_yes": k_yes, "k_no": k_no, "k_series": k.series_ticker, "k_event": k.event_ticker,
        "p_q": p.title, "p_url": p.url, "p_yes": p_yes, "p_no": p_no,
        "gross": arb["gross_edge"], "fees_poly": arb["poly_fee"], "fees_kal": arb["kalshi_fee"],
        "poly_cat": p.fee_cat or "n/a",
        "closes": (p.close_time or k.close_time),
        "score": score, "ok": ok,
    }


def build_all_row(arb: dict, p: Market, k: Market, score: float, inverted: bool, flags: list[str]) -> dict:
    """One line per matched pair (arbitrage or not) for the 'All matches' sheet."""
    p_yes, _ = _yes_no_prices(p)
    k_yes, _ = _yes_no_prices(k)
    net = arb["net_edge"]
    kind = "multi-choice / ladder option" if "option" in (p.kind, k.kind) else (
        "game" if p.kind == "game" else "single question")
    return {"net": None if net == float("-inf") else net, "p_q": p.title, "p_yes": p_yes, "p_url": p.url,
            "k_q": k.title, "k_ticker": k.market_id, "k_yes": k_yes, "score": score, "type": kind,
            "flags": ", ".join(flags) if flags else "OK"}


HEADERS_XLSX = [
    ("#", 5), ("Status", 18), ("Net profit per $1 payout", 14), ("Return on capital", 11),
    (f"Profit on ${CAPITAL:,.0f}", 12), ("% capital → Polymarket", 12), ("% capital → Kalshi", 12),
    ("What to buy", 46),
    ("Kalshi question", 52), ("Kalshi link", 12), ("Kalshi market ticker", 26), ("Kalshi YES bid/ask", 15), ("Kalshi NO bid/ask", 15),
    ("Polymarket question", 52), ("Polymarket link", 14), ("Poly YES bid/ask", 15), ("Poly NO bid/ask", 15),
    ("Gross edge", 10), ("Fees Polymarket", 11), ("Fees Kalshi", 11), ("Poly fee category", 14),
    ("Closes", 12), ("Match score", 9),
]

NOTES = [
    "HOW TO READ THIS FILE",
    "Net profit per $1 payout: guaranteed profit per contract pair AFTER both venues' taker fees (buy one side on each venue, one leg always pays $1).",
    "Return on capital: net profit / total money spent on the pair (prices + fees).",
    "% capital: split your capital between the two legs in these proportions so both legs buy the same number of contracts.",
    "Rows are sorted: 'OK' rows first, then rows flagged CHECK; inside each group the best return on capital first.",
    "CHECK = huge edge (usually the two questions resolve differently), wide bid-ask spread, missing order book, a dust-priced leg (<1c), or a guard warning.",
    "The profit is only guaranteed if BOTH questions resolve identically. contract_guard.py rejects pairs that differ in a number / rank / threshold wording (at least vs above vs exactly) / calendar day / weather station / scope (US vs global) / game sub-market (1H, halftime, race to N); 'WARN' rows passed those checks but share little wording - read both rules.",
    "Kalshi link = the event page (kalshi.com/markets/<series>/<slug>/<event ticker>). The 'Kalshi market ticker' column identifies the exact contract inside that event.",
    "Prices are best bid/ask of the YES token. NO ask = 1 - YES bid, NO bid = 1 - YES ask.",
    "Sheet 'All matches' lists EVERY pair the matcher accepted (arbitrage or not); sheet 'Not matched' lists the closest pairs it rejected and why (including 'guard: ...' rejections).",
    "Multi-choice / ladder markets: the two EVENTS (questions) are matched first, then each option is paired with the SAME option on the other platform (France<->France, $100k rung<->$100k rung). Each pair is an independent Yes/No bet. The INFO flag shows which two events were paired - check they resolve the same way.",
    "Sheet 'Unpaired events' lists every multi-choice / ladder event (either platform) that found no counterpart, with its closest candidate on the other platform and why it was not paired. High similarity + 'no option matched' usually means a naming problem (add an alias); low similarity usually means the event exists on one platform only. Events dropped before matching (past the load limits, non Yes/No outcomes) are NOT listed - see the run log stats.",
    "WARN = accepted but names only partly match (e.g. 'Benjamin Netanyahu' vs 'Netanyahu'), or the guard found low wording overlap: confirm it is the same person/entity/question.",
    "NOT included: order-book depth (size you can actually fill), slippage, USDC/bridging costs, taxes. Gamma's best bid/ask can lag (especially live games): re-check both order books before trading.",
    "ALWAYS read both venues' resolution rules before trading: matching is automatic and can be wrong.",
    "Fee rates: Kalshi 0.07*P*(1-P) rounded up per order; Polymarket category rate * p*(1-p). Verify they are current.",
]


def write_report(rows: list[dict], path: Path, all_rows: list[dict] | None = None,
                 near: list | None = None, unpaired: list | None = None) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:                       # fallback: plain CSV
        import csv
        path = path.with_suffix(".csv")
        with path.open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow([h for h, _ in HEADERS_XLSX])
            for i, r in enumerate(rows, 1):
                w.writerow([i, r["status"], f"{r['net_per_unit']:.4f}", f"{r['roi']:.4%}",
                            f"{r['profit_on_capital']:.2f}", f"{r['poly_pct']:.1%}", f"{r['kal_pct']:.1%}",
                            r["action"], r["k_q"], r["k_url"], r["k_ticker"], r["k_yes"], r["k_no"],
                            r["p_q"], r["p_url"], r["p_yes"], r["p_no"], f"{r['gross']:.4f}",
                            f"{r['fees_poly']:.4f}", f"{r['fees_kal']:.4f}", r["poly_cat"],
                            r["closes"].strftime("%Y-%m-%d") if r["closes"] else "", f"{r['score']:.0f}"])
        return path

    wb = Workbook()
    ws = wb.active
    ws.title = "Arbitrage"
    head_fill = PatternFill("solid", fgColor="1F3864")
    for c, (name, width) in enumerate(HEADERS_XLSX, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 44
    ws.freeze_panes = "C2"

    ok_fill = PatternFill("solid", fgColor="E2F0D9")
    chk_fill = PatternFill("solid", fgColor="FFF2CC")
    wrap = Alignment(wrap_text=True, vertical="top")
    link_font = Font(color="0563C1", underline="single")

    for i, r in enumerate(rows, 1):
        vals = [i, r["status"], r["net_per_unit"], r["roi"], r["profit_on_capital"],
                r["poly_pct"], r["kal_pct"], r["action"],
                r["k_q"], "Open on Kalshi", r["k_ticker"], r["k_yes"], r["k_no"],
                r["p_q"], "Open on Polymarket", r["p_yes"], r["p_no"],
                r["gross"], r["fees_poly"], r["fees_kal"], r["poly_cat"],
                r["closes"].strftime("%Y-%m-%d") if r["closes"] else "", round(r["score"])]
        ws.append(vals)
        row = i + 1
        for c in range(1, len(vals) + 1):
            ws.cell(row=row, column=c).alignment = wrap
        ws.cell(row=row, column=2).fill = ok_fill if r["ok"] else chk_fill
        ws.cell(row=row, column=3).number_format = "$0.0000"
        ws.cell(row=row, column=4).number_format = "0.00%"
        ws.cell(row=row, column=5).number_format = "$#,##0.00"
        ws.cell(row=row, column=6).number_format = "0.0%"
        ws.cell(row=row, column=7).number_format = "0.0%"
        for c in (18, 19, 20):
            ws.cell(row=row, column=c).number_format = "$0.0000"
        for col, url in ((10, r["k_url"]), (15, r["p_url"])):
            if url:
                cell = ws.cell(row=row, column=col)
                cell.hyperlink = url
                cell.font = link_font

    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS_XLSX))}{max(2, len(rows) + 1)}"

    def _sheet(title, headers, data, widths, wrap_cols=()):
        sh = wb.create_sheet(title)
        for c, (h, w) in enumerate(zip(headers, widths), 1):
            cell = sh.cell(row=1, column=c, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            sh.column_dimensions[get_column_letter(c)].width = w
        for row_vals in data:
            sh.append(row_vals)
        for row_cells in sh.iter_rows(min_row=2):
            for cell in row_cells:
                cell.alignment = wrap
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(2, len(data) + 1)}"
        return sh

    if all_rows is not None:
        data = [[i, r["net"], r["type"], r["p_q"], r["p_yes"], r["k_q"], r["k_ticker"], r["k_yes"],
                 round(r["score"]), r["flags"]] for i, r in enumerate(all_rows, 1)]
        sh = _sheet("All matches",
                    ["#", "Net edge per $1 (after fees)", "Type", "Polymarket question", "Poly YES bid/ask",
                     "Kalshi question", "Kalshi market ticker", "Kalshi YES bid/ask", "Match score", "Flags"],
                    data, [5, 14, 16, 52, 15, 52, 26, 15, 9, 40])
        for rr in range(2, len(data) + 2):
            sh.cell(row=rr, column=2).number_format = "$0.0000"
        for rr, r in enumerate(all_rows, 2):
            if r["p_url"]:
                sh.cell(row=rr, column=4).hyperlink = r["p_url"]

    if near is not None:
        data = [[round(sc), why, p.title, k.title, k.market_id] for sc, why, p, k in near]
        _sheet("Not matched (near misses)",
               ["Similarity", "Why rejected", "Polymarket question", "Kalshi question", "Kalshi market ticker"],
               data, [10, 40, 55, 55, 26])

    if unpaired is not None:
        _sheet("Unpaired events",
               ["Platform", "Event (multi-choice / ladder)", "# options", "First options",
                "Closest event on the other platform", "Similarity", "Why not paired", "Link / ticker"],
               unpaired, [11, 50, 9, 45, 50, 10, 52, 30])

    notes = wb.create_sheet("Notes")
    notes.column_dimensions["A"].width = 140
    for i, line in enumerate(NOTES, 1):
        c = notes.cell(row=i, column=1, value=line)
        c.alignment = Alignment(wrap_text=True)
        if i == 1:
            c.font = Font(bold=True, size=13)
    wb.save(path)
    return path
